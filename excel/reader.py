from openpyxl import load_workbook

from config import (
    EXCEL_SHEET,
    CONTROL_NUMBER_COLUMN,
    FIRST_DATA_ROW
)

from models.document import Document

from utils.logger import get_logger

logger = get_logger(__name__)


class ExcelReader:

    def read(self, file_path):

        logger.info(f"Reading workbook: {file_path}")

        workbook = load_workbook(file_path)

        worksheet = workbook[EXCEL_SHEET]

        documents = []

        row = FIRST_DATA_ROW

        while True:

            control_number = worksheet[
                f"{CONTROL_NUMBER_COLUMN}{row}"
            ].value

            if not control_number:
                break

            documents.append(

                Document(

                    row=row,

                    control_number=str(control_number).strip()
                )

            )

            row += 1

        logger.info(
            f"{len(documents)} document(s) loaded."
        )

        return documents